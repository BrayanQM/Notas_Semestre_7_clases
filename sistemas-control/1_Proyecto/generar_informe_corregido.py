#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera la versión corregida del informe con los componentes de INF_EMB."""

import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import generar_informe_math as gen

PROYECTO = "/home/eynor/Documents/5_Class/sistemas-control/1_Proyecto"
SALIDA = os.path.join(PROYECTO, "3_Proyecto_SisContrl_Corregido.docx")
MD = "/tmp/opencode/diseno_math/contenido_final.md"
REF = "/tmp/opencode/diseno_math/referencia_final.docx"
TITULO = (
    "SISTEMA DE MONITOREO DE SIGNOS VITALES BASADO EN ESP32-C3 "
    "PARA SUPERVISIÓN REMOTA DE PERSONAS MAYORES EN LA PAZ"
)


def generar_docx():
    gen.SALIDA = SALIDA
    gen.MD = MD
    gen.REF = REF
    gen.crear_referencia()
    gen.ejecutar([
        "pandoc", MD, "-o", SALIDA,
        "--reference-doc", REF,
        "--from", "markdown",
        "--standalone",
    ])
    gen.postprocesar()


def ajustar_documento():
    doc = Document(SALIDA)

    doc.core_properties.title = TITULO
    doc.core_properties.subject = "Sistemas de Control"
    doc.core_properties.keywords = (
        "ESP32-C3, MAX30102, MAX30205, MLX90614, signos vitales, altitud, IoT"
    )

    # El TOC estático viene del informe base y conserva el título anterior.
    for p in doc.paragraphs:
        if p.style.name.lower() not in ("toc 1", "toc 2", "toc 3"):
            continue
        if not p.text.startswith("6.6 "):
            continue
        for r in p.runs:
            if "Dispositivo portátil" in r.text:
                r.text = r.text.replace(
                    "Dispositivo portátil",
                    "Dispositivo de monitoreo domiciliario",
                )

    for p in doc.paragraphs:
        texto = p.text.strip()

        # Imágenes: centradas y sin sangría de primera línea.
        if p._p.findall(".//" + qn("w:drawing")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            continue

        # Capturas de figuras y tablas con tipografía uniforme.
        if re.match(r"^(Figura|Tabla)\s+\d", texto):
            es_figura = texto.startswith("Figura")
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if es_figura
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(12 if es_figura else 6)
            if es_figura:
                p.paragraph_format.keep_together = True
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(10)
                r.font.italic = True
                if not es_figura:
                    r.font.bold = True

    doc.save(SALIDA)


def convertir_pdf():
    if not shutil_which("soffice"):
        print("ADVERTENCIA: no se encontró soffice; no se generó el PDF.")
        return None
    resultado = subprocess.run(
        [
            "soffice", "--headless", "--convert-to", "pdf",
            "--outdir", PROYECTO, SALIDA,
        ],
        capture_output=True,
        text=True,
    )
    pdf = os.path.join(PROYECTO, "3_Proyecto_SisContrl_Corregido.pdf")
    if resultado.returncode != 0 or not os.path.exists(pdf):
        print(resultado.stdout)
        print(resultado.stderr, file=sys.stderr)
        return None
    return pdf


def shutil_which(nombre):
    from shutil import which
    return which(nombre)


def validar(pdf):
    doc = Document(SALIDA)
    print("DOCX:", SALIDA)
    print("Párrafos:", len(doc.paragraphs))
    print("Tablas:", len(doc.tables))
    print("Imágenes inline:", len(doc.inline_shapes))
    print("Ecuaciones display:", len(doc.element.body.findall(".//" + qn("m:oMathPara"))))
    print("Ecuaciones inline:", len(doc.element.body.findall(".//" + qn("m:oMath"))))
    comentarios = doc.element.body.findall(".//" + qn("w:commentReference"))
    print("Comentarios:", len(comentarios))
    entradas = [
        p for p in doc.paragraphs
        if p.style.name.lower() in ("toc 1", "toc 2", "toc 3") and "\t" in p.text
    ]
    print("Entradas TOC:", len(entradas))
    pendientes = [
        p.text.strip() for p in doc.paragraphs
        if p.text.strip().startswith("PENDIENTE")
    ]
    print("Pendientes:", len(pendientes))
    for texto in pendientes:
        print("  -", texto[:90])
    if pdf and os.path.exists(pdf):
        info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
        paginas = re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE)
        print("PDF:", pdf)
        print("Páginas:", paginas.group(1) if paginas else "?")


def main():
    if not os.path.exists(MD):
        raise SystemExit("No se encontró el Markdown final: " + MD)
    generar_docx()
    ajustar_documento()
    gen.actualizar_indice_desde_pdf()
    pdf = convertir_pdf()
    validar(pdf)


if __name__ == "__main__":
    main()
